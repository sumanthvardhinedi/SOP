"""Read one genuine .xlsx worksheet without evaluating formulas or saving files."""
from collections.abc import Iterator
from io import BytesIO
from zipfile import ZipFile

from defusedxml import ElementTree
from openpyxl import load_workbook
from openpyxl.xml.constants import XLSX

from app.uploads.schemas import SalesValidationError, ValidationIssue

# Bound expansion and worksheet iteration as well as compressed upload size.
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_DATA_ROWS = 50_000
MAX_ARCHIVE_MEMBERS = 1_000


def iter_excel_rows(contents: bytes) -> Iterator[tuple]:
    """Yield cells, including formulas/errors so the service can reject them."""
    workbook = None
    try:
        with ZipFile(BytesIO(contents)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
                raise SalesValidationError(
                    [ValidationIssue(message="Workbook exceeds the expanded file size limit.")], 413
                )
            manifest = ElementTree.fromstring(archive.read("[Content_Types].xml"))
            content_types = [entry.get("ContentType", "") for entry in manifest]
            if XLSX not in content_types or any(
                "macroenabled" in kind.lower() or "vbaproject" in kind.lower()
                for kind in content_types
            ):
                raise SalesValidationError(
                    [ValidationIssue(message="Only standard .xlsx workbooks are supported.")], 415
                )

        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=False, keep_links=False)
        if len(workbook.sheetnames) != 1 or len(workbook.worksheets) != 1:
            raise SalesValidationError(
                [ValidationIssue(message="Workbook must contain exactly one worksheet.")]
            )
        sheet = workbook.worksheets[0]
        # Do not trust producer-supplied dimensions, which can hide actual cells.
        sheet.reset_dimensions()
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            if row_number > MAX_DATA_ROWS + 1:
                raise SalesValidationError(
                    [ValidationIssue(message=f"Workbook exceeds the {MAX_DATA_ROWS} data-row limit.")], 413
                )
            yield cells
    except SalesValidationError:
        raise
    except Exception as exc:
        # Malformed ZIP/XML and invalid Excel structures have many parser error types.
        # Report one safe client error without exposing internals or treating it as valid.
        raise SalesValidationError(
            [ValidationIssue(message="File is not a readable .xlsx workbook.")], 400
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
